import docx
import pandas as pd
from os.path import join as os_join
import os

dataframe_location = ['output','dataframes_pkl']
dataframes = os.listdir(os_join(*dataframe_location))
dataframes_doc = {}
for df in dataframes:
    dataframe_location.append(df)
    dataframes_doc[df[:-4]] = pd.read_pickle(os_join(*dataframe_location))
    del dataframe_location[-1]

df = dataframes_doc['Ergebnisse']
f = os_join(*['output', 'report.docx'])
doc = docx.Document()
header_rows = df.axes[1].nlevels

table = doc.add_table(df.shape[0]+1, 6)#df.shape[1])

# add the header rows.
header = df.columns.tolist()
#for i in range(header_rows):
for j in range(6):#range(df.shape[-1]):
    table.cell(0,j).text = str(header[j][header_rows-1])

# add the rest of the data frame
for i in range(df.shape[0]):
    for j in range(6):
        table.cell(i+1,j).text = str(df.values[i,j])

doc.save(f)


